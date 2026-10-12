"""Load common competition document formats into page/section-aware units."""

import csv
import json
from html.parser import HTMLParser
from pathlib import Path
from typing import Iterable, List

from .cleaning import clean_text, remove_repeated_page_lines
from .models import SourceUnit


SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".html", ".htm", ".json", ".csv", ".docx", ".pptx", ".xlsx"}


class _VisibleHTML(HTMLParser):
    def __init__(self):
        super().__init__()
        self.hidden = 0
        self.parts = []

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "noscript", "svg"}:
            self.hidden += 1
        if tag in {"p", "div", "section", "article", "li", "tr", "h1", "h2", "h3", "h4", "br"}:
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in {"script", "style", "noscript", "svg"} and self.hidden:
            self.hidden -= 1
        if tag in {"p", "div", "section", "article", "li", "tr", "h1", "h2", "h3", "h4"}:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def _relative(path: Path, root: Path) -> str:
    try:
        return path.relative_to(root).as_posix()
    except ValueError:
        return path.name


def _load_pdf(path: Path, source: str) -> List[SourceUnit]:
    try:
        import pdfplumber
    except ImportError as exc:
        raise RuntimeError("PDF support requires pdfplumber") from exc
    units = []
    with pdfplumber.open(path) as pdf:
        total_pages = len(pdf.pages)
        for number, page in enumerate(pdf.pages, 1):
            text = page.extract_text(x_tolerance=2, y_tolerance=3) or ""
            if clean_text(text):
                units.append(SourceUnit(text=text, source=source, page=number, document_type="pdf"))
    if total_pages and len(units) < total_pages / 2:
        print(
            f"Warning: {source} yielded text from only {len(units)}/{total_pages} pages; "
            "it may be scanned and require OCR."
        )
    return units


def _load_docx(path: Path, source: str) -> List[SourceUnit]:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError("DOCX support requires python-docx") from exc
    document = Document(path)
    units, section, paragraphs = [], "", []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            if paragraphs:
                paragraphs.append("")
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style else ""
        if style.startswith("heading"):
            if paragraphs:
                units.append(SourceUnit("\n".join(paragraphs), source, section=section, document_type="docx"))
                paragraphs = []
            section = text
        paragraphs.append(text)
    if paragraphs:
        units.append(SourceUnit("\n".join(paragraphs), source, section=section, document_type="docx"))
    return units


def _load_pptx(path: Path, source: str) -> List[SourceUnit]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError("PPTX support requires python-pptx") from exc
    presentation = Presentation(path)
    units = []
    for number, slide in enumerate(presentation.slides, 1):
        texts = [shape.text for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
        if texts:
            units.append(SourceUnit("\n".join(texts), source, page=number, document_type="pptx"))
    return units


def _load_xlsx(path: Path, source: str) -> List[SourceUnit]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("XLSX support requires openpyxl") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    units = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            values = [str(value).strip() if value is not None else "" for value in row]
            if any(values):
                rows.append(" | ".join(values))
        if rows:
            units.append(SourceUnit("\n".join(rows), source, section=sheet.title, document_type="xlsx"))
    return units


def _json_text(value, prefix="") -> List[str]:
    if isinstance(value, dict):
        lines = []
        for key, child in value.items():
            child_prefix = f"{prefix}.{key}" if prefix else str(key)
            lines.extend(_json_text(child, child_prefix))
        return lines
    if isinstance(value, list):
        lines = []
        for index, child in enumerate(value):
            lines.extend(_json_text(child, f"{prefix}[{index}]"))
        return lines
    return [f"{prefix}: {value}"]


def load_file(path: Path, root: Path) -> List[SourceUnit]:
    source = _relative(path, root)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _load_pdf(path, source)
    if suffix == ".docx":
        return _load_docx(path, source)
    if suffix == ".pptx":
        return _load_pptx(path, source)
    if suffix == ".xlsx":
        return _load_xlsx(path, source)
    if suffix in {".txt", ".md"}:
        return [SourceUnit(path.read_text(encoding="utf-8", errors="replace"), source, document_type=suffix[1:])]
    if suffix in {".html", ".htm"}:
        parser = _VisibleHTML()
        parser.feed(path.read_text(encoding="utf-8", errors="replace"))
        return [SourceUnit("".join(parser.parts), source, document_type="html")]
    if suffix == ".json":
        value = json.loads(path.read_text(encoding="utf-8"))
        return [SourceUnit("\n".join(_json_text(value)), source, document_type="json")]
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", errors="replace", newline="") as handle:
            reader = csv.DictReader(handle)
            rows = [" | ".join(f"{key}: {value}" for key, value in row.items()) for row in reader]
        return [SourceUnit("\n".join(rows), source, document_type="csv")]
    return []


def load_documents(directory: Path) -> List[SourceUnit]:
    if not directory.exists():
        raise FileNotFoundError(f"Documents directory does not exist: {directory}")
    paths = sorted(path for path in directory.rglob("*") if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS)
    if not paths:
        raise RuntimeError(f"No supported documents found in {directory}")
    units = []
    failures = []
    for path in paths:
        try:
            units.extend(load_file(path, directory))
        except Exception as exc:
            failures.append(f"{_relative(path, directory)}: {exc}")
    if not units:
        raise RuntimeError("No text could be extracted. " + "; ".join(failures))
    cleaned = []
    for unit in units:
        text = clean_text(unit.text)
        if text:
            cleaned.append(SourceUnit(text, unit.source, unit.page, unit.section, unit.document_type))
    cleaned = remove_repeated_page_lines(cleaned)
    if failures:
        print("Warnings while loading documents:")
        for failure in failures:
            print(f"- {failure}")
    return cleaned


def document_inventory(directory: Path) -> dict:
    paths = [path for path in directory.rglob("*") if path.is_file()]
    supported = [path for path in paths if path.suffix.lower() in SUPPORTED_EXTENSIONS]
    return {
        "total_files": len(paths),
        "supported_files": len(supported),
        "unsupported": sorted({path.suffix.lower() or "<no extension>" for path in paths if path not in supported}),
        "by_extension": {
            extension: sum(path.suffix.lower() == extension for path in supported)
            for extension in sorted({path.suffix.lower() for path in supported})
        },
    }
